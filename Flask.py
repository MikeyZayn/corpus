import os
import io
import datetime
import logging
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, jsonify
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from pymongo import MongoClient, ASCENDING, TEXT
from bson import ObjectId

# External libs
import pdfplumber
import docx
import unicodedata
import re
from deep_translator import GoogleTranslator, MyMemoryTranslator
from groq import Groq
from dotenv import load_dotenv
import langdetect # For basic language detection
from openpyxl import Workbook # For Excel export

# ---------- Load environment variables ----------
load_dotenv()
SECRET_KEY = os.environ.get("FLASK_SECRET")
MONGO_URI = os.environ.get("MONGO_URI")
GROQ_API_KEY = os.environ.get("GROQ_API_KEY") # Ensure Groq API key is loaded

# ---------- Flask init ----------
app = Flask(__name__)
app.config["SECRET_KEY"] = SECRET_KEY
app.config["MONGO_URI"] = MONGO_URI
app.config['UPLOAD_FOLDER'] = '/tmp' # Using /tmp for temporary uploads

# ---------- Logging init ----------
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# ---------- DB init ----------
if not MONGO_URI:
    raise RuntimeError("MONGO_URI environment variable not set")
if not GROQ_API_KEY:
    raise RuntimeError("GROQ_API_KEY environment variable not set")


client = MongoClient(app.config["MONGO_URI"])
db = client["ZuluEnglishCorpus"]
users = db["users"]
sentences = db["sentences"] # Renamed for consistency
logs = db["logs"]

# Create text index for searching
sentences.create_index(
    [("zulu", TEXT), ("english", TEXT), ("domain", ASCENDING)],
    name="text_index",
    default_language="english"
)
users.create_index([("email", ASCENDING)], unique=True)
users.create_index([("username", ASCENDING)], unique=True)

# ---------- File upload config ----------
ALLOWED_EXT = {"pdf", "docx"}
MAX_FILE_SIZE_MB = 20 # 20MB limit

# ---------- Helpers ----------
def log_action(action, user=None, details=None):
    logs.insert_one({
        "timestamp": datetime.datetime.utcnow(),
        "user_email": user, # Use user_email for consistency
        "action": action,
        "details": details
    })

def login_required(f):
    @wraps(f)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            flash("Please log in to continue.", "warning")
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapped

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXT

# ---------- Extractor ----------
def extract_text(file_stream, filename):
    temp_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file_stream.save(temp_path)

    text = []
    if filename.endswith(".pdf"):
        with pdfplumber.open(temp_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
    elif filename.endswith(".docx"):
        doc = docx.Document(temp_path)
        text = [para.text for para in doc.paragraphs if para.text.strip()]
    else:
        raise ValueError("Unsupported file type")
    
    os.remove(temp_path) # Clean up temporary file
    return "\n".join(text)

def clean_text(text):
    text = re.sub(r'^\s*\d+\s*$', '', text, flags=re.MULTILINE) # Remove lines with only numbers
    text = re.sub(r'(OKUQUKETHWE|TABLE OF CONTENTS|CONTENTS|IZITHUNGO|INTRODUCTION|IMIQULU).*', '', text, flags=re.IGNORECASE) # Remove common TOC/section headers
    text = re.sub(r'[\t\r\f\v]+', ' ', text) # Replace tabs/other whitespace with space
    text = re.sub(r'\n\s*\n', '\n', text) # Remove excessive newlines
    text = re.sub(r'\s{2,}', ' ', text) # Replace multiple spaces with single space
    return text.strip()

def normalize_text(text):
    return unicodedata.normalize("NFKD", text)

def split_sentences(text):
    # Improved sentence splitting for both English and Zulu, handling common punctuation
    sentences = re.split(r'(?<=[.!?])\s+|\n', text)
    sentences = [s.strip() for s in sentences if s.strip()]
    # Filter out very short or non-alphabetic sentences that might be artifacts
    sentences = [s for s in sentences if len(s) > 10 and re.search(r'[a-zA-Z0-9]', s)]
    return sentences

# ---------- Translator + Classifier ----------
def get_language_from_text(text):
    try:
        # Detect dominant language for the overall text
        lang_code = langdetect.detect(text)
        if lang_code == "zu":
            return "zulu"
        elif lang_code == "en":
            return "english"
        else:
            return "other"
    except Exception as e:
        logging.warning(f"Language detection failed: {e}. Defaulting to 'other'.")
        return "other"

def translate_batch(sentences, source_lang, batch_size=50):
    results = []
    
    # Initialize translators once per batch call
    translator_en_zu = GoogleTranslator(source="en", target="zu")
    translator_zu_en = GoogleTranslator(source="zu", target="en")
    
    # Additional languages using MyMemoryTranslator as GoogleTranslator has limits for other SA langs
    # Note: MyMemoryTranslator might require API key for higher usage, or has rate limits.
    translator_en_xh = MyMemoryTranslator(source="en", target="xh")
    translator_zu_xh = MyMemoryTranslator(source="zu", target="xh")
    translator_en_nso = MyMemoryTranslator(source="en", target="nso")
    translator_zu_nso = MyMemoryTranslator(source="zu", target="nso")
    translator_en_nd = MyMemoryTranslator(source="en", target="nd")
    translator_zu_nd = MyMemoryTranslator(source="zu", target="nd")
    translator_en_st = MyMemoryTranslator(source="en", target="st")
    translator_zu_st = MyMemoryTranslator(source="zu", target="st")

    for i in range(0, len(sentences), batch_size):
        chunk = sentences[i:i+batch_size]
        
        chunk_translations = {}
        for lang_key in ["zulu", "english", "isiXhosa", "Sepedi", "Ndebele", "Sesotho"]:
            chunk_translations[lang_key] = ["translation_unavailable"] * len(chunk) # Default placeholder

        try:
            if source_lang == "zulu":
                chunk_translations["english"] = translator_zu_en.translate_batch(chunk)
                chunk_translations["isiXhosa"] = translator_zu_xh.translate_batch(chunk)
                chunk_translations["Sepedi"] = translator_zu_nso.translate_batch(chunk)
                chunk_translations["Ndebele"] = translator_zu_nd.translate_batch(chunk)
                chunk_translations["Sesotho"] = translator_zu_st.translate_batch(chunk)
                # For source lang, the "translation" is the sentence itself
                chunk_translations["zulu"] = chunk
            elif source_lang == "english":
                chunk_translations["zulu"] = translator_en_zu.translate_batch(chunk)
                chunk_translations["isiXhosa"] = translator_en_xh.translate_batch(chunk)
                chunk_translations["Sepedi"] = translator_en_nso.translate_batch(chunk)
                chunk_translations["Ndebele"] = translator_en_nd.translate_batch(chunk)
                chunk_translations["Sesotho"] = translator_en_st.translate_batch(chunk)
                # For source lang, the "translation" is the sentence itself
                chunk_translations["english"] = chunk
            else:
                logging.error(f"Unsupported source language for translation: {source_lang}")
                for _ in chunk:
                    results.append({"error": "translation_failed"})
                continue # Skip to next chunk if source_lang is unsupported
            
            # Combine individual translations for each sentence in the chunk
            for j, sentence in enumerate(chunk):
                result = {lang: chunk_translations[lang][j] for lang in chunk_translations}
                results.append(result)

        except Exception as e:
            logging.error(f"Batch translation failed for chunk {i}-{i+len(chunk)}: {e}")
            for _ in chunk:
                results.append({"error": "translation_failed"})
    return results

def classify_batch(sentences, batch_size=20):
    client = Groq(api_key=GROQ_API_KEY)
    results = []
    
    # Define available domains with their English explanations for better classification
    cultural_domains = [
        "Ubuntombi (Maidenhood, the state of being a young, unmarried woman)",
        "Ukushada (Marriage, wedding traditions, marital relationships)",
        "Ukubingelela (Greetings, formal and informal salutations)",
        "Ukuvakashela (Visiting, hospitality, etiquette when visiting)",
        "Izaga & Izisho (Proverbs & Sayings, traditional wisdom and idioms)",
        "Ubudlelwano (Relationships, family ties, friendships, community roles)",
        "Imicimbi (Events, ceremonies, celebrations, gatherings)",
        "Ukunxiba & Ubuhle (Dressing & Beauty, traditional attire, grooming, aesthetics)",
        "Ukudla (Food, traditional cuisine, eating customs, food preparation)",
        "Umculo & Umhubhe (Music & Hymn, traditional songs, dances, spiritual hymns)",
        "Ezemfundo (Education, learning, schooling, traditional knowledge transfer)",
        "Ezempilo (Health, wellness, traditional healing practices, sickness)",
        "Ezomnotho (Economy, livelihood, trade, traditional economic activities)",
        "Ezesizwe (National/Tribal Affairs, governance, leadership, community issues)",
        "Ezemvelo (Environment, nature, traditional understanding of surroundings)"
    ]
    domain_list_str = "\n- ".join(cultural_domains)

    for i in range(0, len(sentences), batch_size):
        chunk = sentences[i:i+batch_size]
        prompt = f"""
        Classify each sentence into one of the provided cultural domains.
        For each sentence, output ONLY the domain name (the part before the parenthesis, e.g., Ubuntombi).
        If a sentence does not clearly fit any domain, categorize it as 'Okujwayelekile' (General).

        Example Format:
        1. Sentence one -> DomainName
        2. Sentence two -> AnotherDomainName

        Sentences to classify:
        """
        for idx, sentence in enumerate(chunk, start=1):
            prompt += f"{idx}. {sentence}\n"
        
        prompt += f"\nAvailable Domains:\n- {domain_list_str}\n- Okujwayelekile (General, miscellaneous topics)"

        try:
            completion = client.chat.completions.create(
                model="llama3-8b-8192", # Using a smaller, faster model
                messages=[{"role": "user", "content": prompt}],
                temperature=0, # Deterministic output
                max_tokens=500 # Adjust as needed for batch size
            )
            output_lines = completion.choices[0].message.content.strip().splitlines()
            
            chunk_results = []
            for line in output_lines:
                # Extract domain name, handling potential numbering or extra text from LLM
                match = re.search(r'->\s*([a-zA-Z\s&]+?)(?:\s*\(|$)', line)
                if match:
                    domain = match.group(1).strip()
                    # Map the extracted domain to the simpler name
                    if "Ubuntombi" in domain: domain = "Ubuntombi"
                    elif "Ukushada" in domain: domain = "Ukushada"
                    elif "Ukubingelela" in domain: domain = "Ukubingelela"
                    elif "Ukuvakashela" in domain: domain = "Ukuvakashela"
                    elif "Izaga & Izisho" in domain: domain = "Izaga & Izisho"
                    elif "Ubudlelwano" in domain: domain = "Ubudlelwano"
                    elif "Imicimbi" in domain: domain = "Imicimbi"
                    elif "Ukunxiba & Ubuhle" in domain: domain = "Ukunxiba & Ubuhle"
                    elif "Ukudla" in domain: domain = "Ukudla"
                    elif "Umculo & Umhubhe" in domain: domain = "Umculo & Umhubhe"
                    elif "Ezemfundo" in domain: domain = "Ezemfundo"
                    elif "Ezempilo" in domain: domain = "Ezempilo"
                    elif "Ezomnotho" in domain: domain = "Ezomnotho"
                    elif "Ezesizwe" in domain: domain = "Ezesizwe"
                    elif "Ezemvelo" in domain: domain = "Ezemvelo"
                    elif "Okujwayelekile" in domain: domain = "Okujwayelekile"
                    else: domain = "Uncategorized" # Fallback if LLM gives an unexpected domain
                    chunk_results.append(domain)
                else:
                    chunk_results.append("Uncategorized") # Default if parsing fails
            
            # Ensure results match chunk size, fill with 'Uncategorized' if LLM missed some
            while len(chunk_results) < len(chunk):
                chunk_results.append("Uncategorized")
            results.extend(chunk_results)

        except Exception as e:
            logging.error(f"Batch classification failed for chunk {i}-{i+len(chunk)}: {e}")
            results.extend(["Classification_Failed"] * len(chunk))
    return results

def process_file(file_stream, filename, source_lang_override=None):
    raw_text = extract_text(file_stream, filename)
    
    # Attempt to detect language if not overridden
    detected_lang = get_language_from_text(raw_text)
    source_lang = source_lang_override if source_lang_override in ["zulu", "english"] else detected_lang

    if source_lang == "other":
        logging.warning(f"Detected language '{detected_lang}' for {filename}. Processing may be suboptimal.")
        # Decide if you want to flash a warning or proceed with a default
        # For now, we'll proceed assuming the user's selected language or a default.

    cleaned = clean_text(raw_text)
    normalized = normalize_text(cleaned)
    sentences_list = split_sentences(normalized) # Renamed to avoid conflict with sentences collection

    if not sentences_list:
        return []

    translations = translate_batch(sentences_list, source_lang)
    domains = classify_batch(sentences_list)

    structured_entries = []
    for s_idx, s in enumerate(sentences_list):
        t = translations[s_idx]
        d = domains[s_idx]
        
        # Ensure we always have zulu and english fields
        zulu_val = s if source_lang == "zulu" else t.get("zulu", "Translation_Missing")
        english_val = s if source_lang == "english" else t.get("english", "Translation_Missing")

        # Fallback for classification issues
        domain_val = d if d not in ["classification_failed", "Uncategorized"] else "Okujwayelekile"
        
        entry = {
            "zulu": zulu_val,
            "english": english_val,
            "translations": t, # Store all available translations
            "domain": domain_val,
            "tags": [], # Tags for future manual addition
            "original_sentence": s, # Store original sentence for reference
            "source_language_detected": detected_lang, # The language detected from the file
            "source_language_used_for_processing": source_lang, # The language used for translation logic
            "origin_filename": filename,
            "created_at": datetime.datetime.utcnow(),
            "added_by": session.get("username", "anonymous")
        }
        structured_entries.append(entry)
    return structured_entries


# ---------- Routes ----------
@app.route("/")
def index():
    # If user is logged in, redirect to dashboard
    if "user_id" in session:
        return redirect(url_for("dashboard"))
    
    query = request.args.get("q")
    domain_filter = request.args.get("domain")
    results = []

    search_criteria = {}
    if query:
        search_criteria["$text"] = {"$search": query}
    if domain_filter and domain_filter != "all":
        search_criteria["domain"] = domain_filter

    if search_criteria:
        cursor = sentences.find(search_criteria).limit(50)
        results = list(cursor)

    return render_template("index.html", query=query, results=results, selected_domain=domain_filter)

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form["username"].strip()
        email = request.form["email"].strip().lower()
        password = request.form["password"]

        if users.find_one({"$or": [{"email": email}, {"username": username}]}):
            flash("Email or username already registered.", "warning")
            return redirect(url_for("register"))

        users.insert_one({
            "username": username,
            "email": email,
            "password": generate_password_hash(password),
            "created_at": datetime.datetime.utcnow()
        })
        log_action("user_registered", user=email)
        flash("Account created successfully! Please log in.", "success")
        return redirect(url_for("login"))

    return render_template("register.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        identifier = request.form["username"].strip() # Can be username or email
        password = request.form["password"]

        user = users.find_one({"$or": [{"email": identifier}, {"username": identifier}]})
        if not user:
            flash("Invalid username or email.", "danger")
            return redirect(url_for("login"))

        if check_password_hash(user["password"], password):
            session["user_id"] = str(user["_id"])
            session["user_email"] = user["email"]
            session["username"] = user.get("username", user["email"]) # Fallback to email if username not set
            log_action("user_logged_in", user=user["email"])
            flash(f"Welcome back, {session['username']}!", "success")
            return redirect(url_for("dashboard")) # Redirect to dashboard after login
        else:
            flash("Incorrect password.", "danger")
            return redirect(url_for("login"))

    return render_template("login.html")

@app.route("/logout")
def logout():
    user_email = session.get("user_email")
    session.clear()
    if user_email:
        log_action("user_logged_out", user=user_email)
    flash("You have been logged out.", "info")
    return redirect(url_for("index"))

@app.route("/profile")
@login_required
def profile():
    user = users.find_one({"_id": ObjectId(session["user_id"])})
    if not user:
        flash("User profile not found.", "danger")
        session.clear()
        return redirect(url_for("login"))
    return render_template("profile.html", user=user)

@app.route("/dashboard")
@login_required
def dashboard():
    query = request.args.get("q")
    domain_filter = request.args.get("domain")
    results = []

    search_criteria = {}
    if query:
        search_criteria["$text"] = {"$search": query}
    if domain_filter and domain_filter != "all":
        search_criteria["domain"] = domain_filter

    if search_criteria:
        cursor = sentences.find(search_criteria).limit(50)
        results = list(cursor)
    else:
        # Show recent entries on dashboard by default if no search
        results = list(sentences.find().sort("created_at", -1).limit(20))

    # All unique domains for filter dropdown
    all_domains = sorted(sentences.distinct("domain"))

    return render_template("dashboard.html", 
                           query=query, 
                           results=results, 
                           selected_domain=domain_filter,
                           domains=all_domains)

@app.route("/download")
@login_required
def download_corpus():
    # Determine format from request, default to TXT if not specified or invalid
    export_format = request.args.get("format", "txt").lower()

    if export_format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Corpus Data"

        headers = ["_id", "Zulu", "English", "Domain", "Original Sentence", "Source Language Detected", "Source Language Used", "Origin Filename", "Added By", "Created At"]
        # Add dynamic translation language headers
        sample_entry = sentences.find_one()
        if sample_entry and "translations" in sample_entry:
            for lang in sorted(sample_entry["translations"].keys()):
                if lang not in ["zulu", "english"]: # Avoid duplication
                    headers.append(f"Translation ({lang})")
        
        ws.append(headers)

        for entry in sentences.find():
            row_data = [
                str(entry.get("_id", "")),
                entry.get("zulu", ""),
                entry.get("english", ""),
                entry.get("domain", ""),
                entry.get("original_sentence", ""),
                entry.get("source_language_detected", ""),
                entry.get("source_language_used_for_processing", ""),
                entry.get("origin_filename", ""),
                entry.get("added_by", ""),
                entry.get("created_at", "").isoformat() if entry.get("created_at") else ""
            ]
            if "translations" in entry:
                for lang in sorted(sample_entry["translations"].keys()):
                    if lang not in ["zulu", "english"]:
                        row_data.append(entry["translations"].get(lang, ""))
            
            ws.append(row_data)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        log_action("download_corpus", user=session.get("user_email"), details={"format": "xlsx"})
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="corpus_dump.xlsx"
        )

    else: # Default to TXT
        buffer = io.StringIO()
        for entry in sentences.find():
            line = f"Zulu: {entry.get('zulu', 'N/A')}\nEnglish: {entry.get('english', 'N/A')}\n"
            if "translations" in entry:
                for lang, val in entry["translations"].items():
                    line += f"{lang.capitalize()}: {val}\n"
            line += f"Domain: {entry.get('domain', 'N/A')}\n"
            line += f"Original Sentence: {entry.get('original_sentence', 'N/A')}\n"
            line += f"Source Detected: {entry.get('source_language_detected', 'N/A')}\n"
            line += f"Origin Filename: {entry.get('origin_filename', 'N/A')}\n"
            line += f"Added By: {entry.get('added_by', 'Anonymous')}\n"
            line += f"Created At: {entry.get('created_at', 'N/A')}\n"
            line += "-" * 40 + "\n\n"
            buffer.write(line)

        buffer.seek(0)
        log_action("download_corpus", user=session.get("user_email"), details={"format": "txt"})
        return send_file(
            io.BytesIO(buffer.getvalue().encode("utf-8")),
            mimetype="text/plain",
            as_attachment=True,
            download_name="corpus_dump.txt"
        )

@app.route("/result/<id>")
def result_detail(id):
    try:
        entry = sentences.find_one({"_id": ObjectId(id)})
        if not entry:
            flash("Result not found.", "warning")
            return redirect(url_for("index"))
    except Exception as e:
        flash(f"Invalid result ID: {e}", "danger")
        return redirect(url_for("index"))

    # Placeholder for future NLP features
    # entry['frequency'] = "Not yet implemented"
    # entry['similar_phrases'] = ["Placeholder 1", "Placeholder 2"]
    # entry['commonly_paired_words'] = ["word A", "word B"]

    return render_template("result_detail.html", entry=entry)

@app.route("/upload", methods=["GET", "POST"])
@login_required
def upload():
    if request.method == "POST":
        if "file" not in request.files:
            flash("No file part in the request.", "danger")
            return redirect(request.url)

        file = request.files["file"]
        if file.filename == "":
            flash("No selected file.", "warning")
            return redirect(request.url)

        if not allowed_file(file.filename):
            flash("Unsupported file type. Only PDF and DOCX are allowed.", "danger")
            return redirect(request.url)

        # Check file size before saving
        file.seek(0, os.SEEK_END)
        file_size = file.tell()
        file.seek(0) # Reset file pointer
        if file_size > MAX_FILE_SIZE_MB * 1024 * 1024:
            flash(f"File size exceeds the {MAX_FILE_SIZE_MB}MB limit.", "danger")
            return redirect(request.url)

        filename = secure_filename(file.filename)
        
        try:
            # User can specify the expected source language, otherwise it will be detected
            source_lang_override = request.form.get("language", "auto").lower()
            if source_lang_override == "auto":
                source_lang_override = None # Let process_file detect

            entries_batch = process_file(file.stream, filename, source_lang_override)

            if entries_batch:
                sentences.insert_many(entries_batch)
                log_action("upload_and_process", user=session.get("user_email"),
                    details={"filename": filename, "entries_count": len(entries_batch), "source_lang_override": source_lang_override})
                flash(f"Successfully uploaded '{filename}' and added {len(entries_batch)} entries to the corpus.", "success")
            else:
                flash(f"No usable entries could be extracted from '{filename}'.", "warning")

        except ValueError as ve:
            log_action("upload_error", user=session.get("user_email"), details={"filename": filename, "error": str(ve)})
            flash(f"File processing error: {ve}", "danger")
        except Exception as e:
            log_action("upload_error", user=session.get("user_email"), details={"filename": filename, "error": str(e)})
            flash(f"An unexpected error occurred during processing: {e}", "danger")

        return redirect(url_for("dashboard"))

    return render_template("upload.html")

@app.route("/healthz")
def healthz():
    return "ok", 200

@app.route("/test")
def test_connection():
    try:
        doc = sentences.find_one()  # get one document from Atlas
        if not doc:
            return {"status": "ok", "message": "Connected, but no documents found."}
        return {
            "status": "ok",
            "sample": {
                "id": str(doc.get("_id")),
                "zulu": doc.get("zulu"),
                "english": doc.get("english"),
                "domain": doc.get("domain")
            }
        }
    except Exception as e:
        return {"status": "error", "message": str(e)}

@app.context_processor
def inject_globals():
    """Injects global variables like language options into all templates."""
    # Define a consistent list of domains with English descriptions for display
    display_domains = [
        ("Ubuntombi", "Maidenhood"), ("Ukushada", "Marriage"),
        ("Ukubingelela", "Greetings"), ("Ukuvakashela", "Visiting"),
        ("Izaga & Izisho", "Proverbs & Sayings"), ("Ubudlelwano", "Relationships"),
        ("Imicimbi", "Events"), ("Ukunxiba & Ubuhle", "Dressing & Beauty"),
        ("Ukudla", "Food"), ("Umculo & Umhubhe", "Music & Hymn"),
        ("Ezemfundo", "Education"), ("Ezempilo", "Health"),
        ("Ezomnotho", "Economy"), ("Ezesizwe", "National Affairs"),
        ("Ezemvelo", "Environment"), ("Okujwayelekile", "General")
    ]
    return dict(display_domains=display_domains)


if __name__ == "__main__":
    app.run(
        debug=True,
        use_reloader=False,  # ✅ prevents WinError 10038 on some systems
        host="0.0.0.0",
        port=int(os.environ.get("PORT", 5000))
    )